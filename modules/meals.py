import streamlit as st
from utils.db import supabase
import google.generativeai as genai
import docx
import io

dias_semana = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]

def generate_word_document(lista_texto):
    doc = docx.Document()
    doc.add_heading('Lista de la Compra Semanal', 0)
    doc.add_paragraph(lista_texto)
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def render_ui():
    st.header("Menú Semanal 🍽️")
    household_id = str(st.session_state.household_id)
    
    # 1. Cargar datos existentes
    res = supabase.table("weekly_meals").select("*").eq("household_id", household_id).execute()
    menu_db = {f"{m['day_of_week']}_{m['meal_type']}": m['dish_name'] for m in res.data}
    
    # 2. Interfaz del Menú
    with st.form("menu_form"):
        col_comida, col_cena = st.columns(2)
        
        nuevos_platos = {}
        
        with col_comida:
            st.subheader("Comidas")
            for dia in dias_semana:
                key = f"{dia}_Comida"
                nuevos_platos[key] = st.text_input(dia, value=menu_db.get(key, ""), key=key)
                
        with col_cena:
            st.subheader("Cenas")
            for dia in dias_semana:
                key = f"{dia}_Cena"
                nuevos_platos[key] = st.text_input(dia, value=menu_db.get(key, ""), key=key)
                
        if st.form_submit_button("Guardar Menú", use_container_width=True):
            data_to_upsert = []
            for dia in dias_semana:
                for tipo in ["Comida", "Cena"]:
                    data_to_upsert.append({
                        "household_id": household_id,
                        "day_of_week": dia,
                        "meal_type": tipo,
                        "dish_name": nuevos_platos[f"{dia}_{tipo}"]
                    })
            supabase.table("weekly_meals").upsert(data_to_upsert, on_conflict="household_id, day_of_week, meal_type").execute()
            st.success("Menú actualizado")
            st.rerun()

    st.divider()

    # 3. Generación de Lista con IA
    if st.button("🛒 Extraer Lista de la Compra con IA", type="primary", use_container_width=True):
        if "GEMINI_API_KEY" not in st.secrets:
            st.error("Falta configurar GEMINI_API_KEY en secrets.toml")
            return
            
        platos_lista = [v for k, v in menu_db.items() if v.strip()]
        if not platos_lista:
            st.warning("El menú está vacío. Añade platos primero.")
            return
            
        texto_menu = "\n".join(platos_lista)
        
        with st.spinner("Analizando recetas y organizando pasillos..."):
            try:
                genai.configure(api_key=st.secrets["GEMINI_API_KEY"])
                model = genai.GenerativeModel('gemini-2.5-flash')
                
                prompt = f"""
                A partir de la siguiente lista de platos, extrae los ingredientes necesarios para hacer la compra.
                Agrúpalos estrictamente por secciones de supermercado (Frutería, Carnicería, Pescadería, Lácteos, Despensa/Congelados, Limpieza).
                Devuelve únicamente la lista con viñetas claras. No añadas introducciones ni conclusiones.
                Platos:
                {texto_menu}
                """
                
                response = model.generate_content(prompt)
                lista_final = response.text
                
                # Crear Word y mostrar botón de descarga
                word_file = generate_word_document(lista_final)
                
                st.success("¡Lista generada!")
                st.download_button(
                    label="📄 Descargar Lista en Word",
                    data=word_file,
                    file_name="Lista_Compra.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
                
                with st.expander("Ver lista previa", expanded=True):
                    st.write(lista_final)
                    
            except Exception as e:
                st.error(f"Error al conectar con la IA: {e}")
